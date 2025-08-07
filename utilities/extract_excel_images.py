import os
import shutil
import zipfile

from PIL import Image

file_path = "/Users/samarsingh/Library/Containers/com.apple.mail/Data/Library/Mail Downloads/4C7D8FC2-7A28-468A-8CE3-EF6B878FA056/ML ENTERPRISES_VOGUE COLLECTIONS 1032.xlsx"

# Method 1: Try using openpyxl to extract images
try:
    from openpyxl import load_workbook

    wb = load_workbook(file_path)
    print("=== CHECKING FOR IMAGES IN EXCEL ===")

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        if hasattr(sheet, "_images") and sheet._images:
            print(f"Sheet '{sheet_name}' has {len(sheet._images)} images")
            for idx, img in enumerate(sheet._images):
                print(f"  - Image {idx+1}: anchor at {img.anchor}")
        else:
            # Check for drawings
            if hasattr(sheet, "_drawing") and sheet._drawing:
                print(f"Sheet '{sheet_name}' has drawings")

    # Check for any drawing parts
    if hasattr(wb, "_drawings") and wb._drawings:
        print(f"Workbook has {len(wb._drawings)} drawing objects")

except Exception as e:
    print(f"Error with openpyxl method: {e}")

# Method 2: Extract images by unzipping Excel file
print("\n=== EXTRACTING IMAGES FROM EXCEL ZIP ===")
temp_dir = "/tmp/excel_extract"
os.makedirs(temp_dir, exist_ok=True)

try:
    # Excel files are actually zip files
    with zipfile.ZipFile(file_path, "r") as zip_ref:
        # List all files in the zip
        file_list = zip_ref.namelist()

        # Look for images in xl/media folder
        image_files = [f for f in file_list if f.startswith("xl/media/")]

        if image_files:
            print(f"Found {len(image_files)} images in Excel file:")
            for img_file in image_files:
                print(f"  - {img_file}")

                # Extract the image
                zip_ref.extract(img_file, temp_dir)

                # Open and get image info
                img_path = os.path.join(temp_dir, img_file)
                with Image.open(img_path) as img:
                    print(
                        f"    Size: {img.size}, Format: {img.format}, Mode: {img.mode}"
                    )

                    # Save to current directory for inspection
                    output_name = f"extracted_{os.path.basename(img_file)}"
                    img.save(output_name)
                    print(f"    Saved as: {output_name}")
        else:
            print("No images found in Excel file")

        # Also check for drawing relationships
        drawing_files = [f for f in file_list if "drawing" in f]
        if drawing_files:
            print(f"\nFound {len(drawing_files)} drawing-related files:")
            for f in drawing_files:
                print(f"  - {f}")

except Exception as e:
    print(f"Error extracting images: {e}")
finally:
    # Cleanup
    if os.path.exists(temp_dir):
        shutil.rmtree(temp_dir)

print("\n=== CHECKING OUR SYSTEM'S CAPABILITY ===")
print("Our system can handle image attachments through:")
print("1. Direct image files (.jpg, .png, etc.) - ✅ Fully supported")
print("2. Images embedded in Excel - ❌ Not currently implemented")
print("3. Images in PDF files - ❌ Not currently implemented (OCR planned)")
print("\nFor embedded Excel images, we would need to:")
print("- Extract images from Excel using openpyxl or by unzipping")
print("- Process each image with Qwen2.5VL for analysis")
print("- Store in ChromaDB with CLIP embeddings")
